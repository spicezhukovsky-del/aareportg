import matplotlib.pyplot as plt
import matplotlib.ticker as ticker
import io

from analyzer.math_def import percent_complet


from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH


def create_table(document, headers, rows, style='Table Grid'):
    cols_number = len(headers)

    table =  document.add_table(rows=1, cols=cols_number)
    table.style = style

    hdr_cells = table.rows[0].cells
    for i in range(cols_number):
        hdr_cells[i].text = headers[i]

    for row in rows:
        row_cells = table.add_row().cells
        for i in range(cols_number):
            row_cells[i].text = str(row[i])

    return table


def create_schedule(name_doc, list_data_x, list_data_y, name_schedule, name_x, name_y):
    fig_sh, ax_sh = plt.subplots(figsize=(7, 2.36))
    ax_sh.plot(list_data_x, list_data_y)
    ax_sh.set_title(name_schedule)
    ax_sh.set_ylabel(name_y)
    ax_sh.set_xlabel(name_x)

    if len(list_data_x) < 20:
        ax_sh.xaxis.set_major_locator(ticker.MultipleLocator(1))
    elif 50 > len(list_data_x) > 20:
        ax_sh.xaxis.set_major_locator(ticker.MultipleLocator(5))
    elif 100 > len(list_data_x) > 50:
        ax_sh.xaxis.set_major_locator(ticker.MultipleLocator(10))
    elif len(list_data_x) > 100:
        ax_sh.xaxis.set_major_locator(ticker.MultipleLocator(20))

    schedule_sh = io.BytesIO()
    plt.savefig(schedule_sh, format='png', dpi=300, bbox_inches='tight')
    schedule_sh.seek(0)
    plt.close(fig_sh)

    return name_doc.add_picture(schedule_sh, width=Cm(18), height=Cm(6))



def over_coach_block(name_doc, dict_data, rank_dict, name_school):
    for name_coach, item in dict_data.items():
        coach_name = name_coach
        # print(f"Имя тренера: {coach_name}")
        total_look = dict_data[name_coach]['general_info']['count_sum']
        # print(f"Кол-во просмотров: {total_look} раз")

        quest = dict_data[name_coach]['questions']
        dict_data[name_coach]['questions'] = percent_complet(total_look, quest)

        mean_grade = dict_data[name_coach]['general_info']['mean_questions']
        # print(f'Средний балл: {mean_grade}')
        best_bad_qurstions = dict_data[name_coach]['questions']

        len_list_rank = len(list(rank_dict.keys()))
        positions_coach = rank_dict[name_coach]['position']

        paragr_name = name_doc.add_paragraph()
        run = paragr_name.add_run(f"{coach_name}\n")
        run.bold = True
        run.font.size = Pt(16)

        bade_text = paragr_name.add_run(f"В этом месяце просмотренно {total_look} тренеровки, получил средний балл {mean_grade}. В этом месяце данный тренер занял {positions_coach} место в рейтенге состоящем из {len_list_rank} мест.")
        bade_text.font.size = Pt(14)


        list_data_in_Y = []
        for data in dict_data[name_coach]['grade']:
            list_data_in_Y.append(data[1])
        list_data_in_X = list(num for num in range(1, len(list_data_in_Y) + 1))

        if len(list_data_in_X) > 9:
            create_schedule(name_doc, list_data_in_X, list_data_in_Y, 'Динамика оценок', 'Просмотренное УТЗ', 'Оценка')

        bad_questions = name_doc.add_paragraph()
        first_run = bad_questions.add_run("Вопросы с самым плохим процентом выполнением:\n")
        first_run.bold = True
        first_run.font.size = Pt(14)

        for question, grade in best_bad_qurstions.items():
            if grade <= 15:
                run = bad_questions.add_run(f"\t{question}: выполнение")
                run_2 = bad_questions.add_run(f" {grade}%\n")
                run_2.bold = True



def sh_block(name_doc, data_dict_one_sh, current_month_rank_dict, name_sh, id_sh, dict_sh_name, past_month_rank_dict = None):


    positions_sh_this_month = current_month_rank_dict[id_sh]['position']

    if past_month_rank_dict is not None:
        positions_sh_past_month = past_month_rank_dict[id_sh]['position']
        dinamick_rank = positions_sh_past_month - positions_sh_this_month
        if dinamick_rank < 0:
            dinamik_text = f"Школа опустилась в рейтенге на {abs(dinamick_rank)} позиций"
        elif dinamick_rank > 0:
            dinamik_text = f"Школа поднялась в рейтенге на {abs(dinamick_rank)} позиций"
        else:
            dinamik_text = "Школа не изменило свое положение в рейтенге"


    mean_grade = data_dict_one_sh['general_info']["mean_questions"]

    count_look = data_dict_one_sh['general_info']["count_sum"]

    quest = data_dict_one_sh['questions']
    data_dict_one_sh['questions'] = percent_complet(count_look, quest)

    len_list_rank = len(list(current_month_rank_dict.keys()))

    positions_this_sh = current_month_rank_dict[id_sh]['position']
    list_in_rank = list(current_month_rank_dict.keys())
    index_positions_this_sh = positions_this_sh - 1
    if positions_this_sh == 1:
        positions_up_sh = None
        text_up_sh = "Школа занимает первое местов общем рейтенге мониторинга."
        positions_down_sh = list_in_rank[index_positions_this_sh + 1]
        text_down_sh = f"{positions_this_sh + 1}) {dict_sh_name[positions_down_sh]}"
    elif positions_this_sh == len_list_rank:
        positions_up_sh = list_in_rank[index_positions_this_sh - 1]
        text_up_sh = f"{positions_this_sh - 1}) {dict_sh_name[positions_up_sh]}"
        positions_down_sh = None
        text_down_sh = "Школа занимает полследенее место в общем рейтенге мониторинга."
    else:
        positions_up_sh = list_in_rank[index_positions_this_sh - 1]
        text_up_sh = f"{positions_this_sh - 1}) {dict_sh_name[positions_up_sh]}"
        positions_down_sh = list_in_rank[index_positions_this_sh + 1]
        text_down_sh = f"{positions_this_sh + 1}) {dict_sh_name[positions_down_sh]}"


    paragr_name = name_doc.add_paragraph()
    run = paragr_name.add_run(f"{name_sh}")
    run.bold = True
    run.font.size = Pt(18)
    run.alignment = WD_ALIGN_PARAGRAPH.CENTER


    bade_text = name_doc.add_paragraph()
    if past_month_rank_dict is not None:
        bade_text_run = bade_text.add_run(f"В этом месяце просмотренно {count_look} тренеровки, средний балл {mean_grade}. В общем рейтенге Школа занимает {positions_this_sh} место. По сравнению с прошлым месяцем {dinamik_text}.")
        bade_text_run.font.size = Pt(14)
        bade_text.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    else:
        bade_text_run = bade_text.add_run(f"В этом месяце просмотренно {count_look} тренеровки, средний балл {mean_grade}. В общем рейтенге Школа занимает {positions_this_sh} место.")
        bade_text_run.font.size = Pt(14)
        bade_text.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY



    rank_text_sh = name_doc.add_paragraph()

    rank_up_sh = rank_text_sh.add_run(f"{text_up_sh}\n")
    rank_up_sh.font.size = Pt(14)

    rank_list = rank_text_sh.add_run(f"{positions_this_sh}) {name_sh}\n")
    rank_list.bold = True
    rank_list.font.size = Pt(18)


    rank_down_sh = rank_text_sh.add_run(f"{text_down_sh}")
    rank_down_sh.font.size = Pt(14)

    list_data_in_Y = []
    for data in data_dict_one_sh['general_info']['over_look']:
        list_data_in_Y.append(data[1])
    list_data_in_X = list(num for num in range(1, len(list_data_in_Y) + 1))

    create_schedule(name_doc, list_data_in_X, list_data_in_Y, 'Динамика оценок', 'Просмотренное УТЗ', 'Оценка')


    bad_questions = name_doc.add_paragraph()
    first_run = bad_questions.add_run("Процент выполнения вопросов:\n")
    first_run.bold = True
    first_run.font.size = Pt(14)

    rang_id = 1
    for question, grade in data_dict_one_sh['questions'].items():
        run = bad_questions.add_run(f"{rang_id}) {question}: выполнение")
        run_2 = bad_questions.add_run(f" {grade}%\n")
        run_2.bold = True
        rang_id += 1
